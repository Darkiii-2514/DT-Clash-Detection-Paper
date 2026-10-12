"""Build the Word manuscript (and highlights) from paper/manuscript.md.

    .venv\\Scripts\\python paper\\tools\\build_docx.py [--pdf]

1. reference.docx: Times New Roman 12 pt, double spacing, A4, 2.5 cm margins (review format)
2. pandoc: numbered sections, citeproc with the Elsevier numeric style (Automation in
   Construction), TeX math -> native Word equations, figures from paper/figures
3. python-docx post-processing: continuous line numbers (journal requirement for review),
   centred page numbers, readable table borders
4. --pdf: export a PDF through Microsoft Word for visual checking
"""
from __future__ import annotations

import argparse
import re
import shutil
import subprocess
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

PAPER = Path(__file__).resolve().parents[1]
PANDOC = shutil.which("pandoc") or str(Path.home() / "AppData/Local/Pandoc/pandoc.exe")
REF = PAPER / "style" / "reference.docx"


def _force_font(style, name="Times New Roman"):
    """Set explicit fonts and drop theme-font attributes (they override explicit ones)."""
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in list(rfonts.attrib):
        if attr.endswith("Theme"):
            del rfonts.attrib[attr]
    for k in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(k), name)


def make_reference():
    # write straight to the file: piping binary through stdout on Windows corrupts the zip
    subprocess.run([PANDOC, "-o", str(REF), "--print-default-data-file", "reference.docx"], check=True)
    d = Document(REF)
    st = d.styles
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Abstract", "Bibliography", "Subtitle", "Author"):
        if name in [s.name for s in st]:
            s = st[name]
            s.font.name = "Times New Roman"
            s.font.size = Pt(12)
            s.font.color.rgb = RGBColor(0, 0, 0)
            _force_font(s)
            pf = s.paragraph_format
            pf.line_spacing = 2.0
            pf.space_after = Pt(0)
            pf.space_before = Pt(0)
    for lvl, size in ((1, 12), (2, 12), (3, 12)):
        s = st[f"Heading {lvl}"]
        s.font.name = "Times New Roman"
        _force_font(s)
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.italic = lvl == 3
        s.font.color.rgb = RGBColor(0, 0, 0)
        s.paragraph_format.space_before = Pt(12)
        s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.line_spacing = 2.0
    for name in ("Title", "Abstract Title"):
        if name not in [x.name for x in st]:
            continue
        s = st[name]
        s.font.name = "Times New Roman"
        _force_font(s)
        s.font.size = Pt(16)
        s.font.bold = True
        s.font.color.rgb = RGBColor(0, 0, 0)
    for name in ("Image Caption", "Table Caption", "Caption"):
        if name in [s.name for s in st]:
            s = st[name]
            s.font.name = "Times New Roman"
            _force_font(s)
            s.font.size = Pt(10)
            s.font.italic = False
            s.font.color.rgb = RGBColor(0, 0, 0)
            s.paragraph_format.line_spacing = 1.15
            s.paragraph_format.space_after = Pt(10)
    sec = d.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, side, Cm(2.5))
    d.save(REF)


def add_line_numbers_and_page_numbers(path: Path):
    d = Document(path)
    for sec in d.sections:
        sp = sec._sectPr
        ln = sp.find(qn("w:lnNumType"))
        if ln is None:
            ln = OxmlElement("w:lnNumType")
            sp.append(ln)
        ln.set(qn("w:countBy"), "1")
        ln.set(qn("w:restart"), "continuous")
        ln.set(qn("w:distance"), "360")
        p = sec.footer.paragraphs[0] if sec.footer.paragraphs else sec.footer.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run()
        for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
            if tag:
                f = OxmlElement("w:fldChar")
                f.set(qn("w:fldCharType"), tag)
                run._r.append(f)
            else:
                t = OxmlElement("w:instrText")
                t.set(qn("xml:space"), "preserve")
                t.text = text
                run._r.append(t)
    # title block: author names bold, affiliation (italic) and corresponding-author lines regular
    authors = [p for p in d.paragraphs if p.style.name == "Author"]
    for i, par in enumerate(authors[1:]):
        for r in par.runs:
            r.bold = False
            r.italic = i == 0
    # tables: thin horizontal rules (booktabs-like), compact single spacing
    for tbl in d.tables:
        tblPr = tbl._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge, val in (("top", "single"), ("bottom", "single"), ("insideH", "single"),
                          ("left", "nil"), ("right", "nil"), ("insideV", "nil")):
            b = OxmlElement(f"w:{edge}")
            b.set(qn("w:val"), val)
            b.set(qn("w:sz"), "4")
            b.set(qn("w:color"), "808080")
            borders.append(b)
        tblPr.append(borders)
        for row in tbl.rows:
            for cell in row.cells:
                for par in cell.paragraphs:
                    par.paragraph_format.line_spacing = 1.0
                    for r in par.runs:
                        r.font.size = Pt(10)
    d.save(path)


def export_pdf(docx: Path, pdf: Path):
    ps = (f"$w = New-Object -ComObject Word.Application; $w.Visible = $false; "
          f"$d = $w.Documents.Open('{docx}', $false, $true); "
          f"$d.ExportAsFixedFormat('{pdf}', 17); $d.Close($false); $w.Quit()")
    subprocess.run(["powershell", "-NoProfile", "-Command", ps], check=True)


ROADMAP = ("The remainder of the paper is organised as follows. Section 2 reviews related work. Section 3 presents "
           "the framework. Section 4 describes the implementation, the case-study models and the evaluation design. "
           "Section 5 reports the results, Section 6 discusses implications and limitations, and Section 7 concludes.")


def to_asce(src: str) -> str:
    """ASCE journals: author-date references, unnumbered headings (so cross-references name the
    section), Data Availability Statement / Acknowledgments / Supplemental Materials instead of the
    Elsevier declarations, and supplemental files named Video S1, ..."""
    titles, top, sub = {}, 0, 0
    for line in src.splitlines():
        m = re.match(r"^(#{1,2}) (.+?)\s*$", line)
        if not m or "{.unnumbered}" in m.group(2):
            continue
        if len(m.group(1)) == 1:
            top, sub = top + 1, 0
            titles[f"{top}"] = m.group(2)
        else:
            sub += 1
            titles[f"{top}.{sub}"] = m.group(2)
    src = src.replace(ROADMAP,                      # only if the numbered roadmap sentence is present"The remainder of the paper reviews related work, presents the framework, describes the "
                      "implementation, the case-study models and the evaluation design, reports the results, and "
                      "discusses implications and limitations before concluding.")
    src = re.sub(r"\(Section (\d+(?:\.\d+)?)\)", lambda m: f"(see “{titles[m.group(1)]}”)", src)
    src = re.sub(r"Section (\d+(?:\.\d+)?)", lambda m: f"the section “{titles[m.group(1)]}”", src)
    src = src.replace("Supplementary Video 1", "Video S1")
    a, b = src.index("# Declarations {.unnumbered}"), src.index("# References {.unnumbered}")
    src = src[:a] + (PAPER / "endmatter_asce.md").read_text(encoding="utf-8") + "\n" + src[b:]
    src = src.replace("csl: style/elsevier-with-titles.csl", "csl: style/american-society-of-civil-engineers.csl")
    # ASCE reference style omits the publisher of journal articles
    entries = (PAPER / "refs.bib").read_text(encoding="utf-8").split("\n@")
    kept = [re.sub(r"\s*publisher\s*=\s*\{[^{}]*\},?", "", e, flags=re.I)
            if e.lower().lstrip("@").startswith("article") else e for e in entries]
    (PAPER / "refs_asce.bib").write_text("\n@".join(kept), encoding="utf-8")
    src = src.replace("bibliography: refs.bib", "bibliography: refs_asce.bib")
    leftover = re.findall(r"Section \d", src)
    assert not leftover, leftover
    return src


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--pdf", action="store_true")
    ap.add_argument("--rebuild-reference", action="store_true")
    ap.add_argument("--journal", choices=("aic", "asce"), default="aic",
                    help="aic: Elsevier format (manuscript.docx); asce: ASCE format (manuscript_asce.docx)")
    a = ap.parse_args()
    if a.rebuild_reference or not REF.exists():
        make_reference()
    asce = a.journal == "asce"
    out = PAPER / ("manuscript_asce.docx" if asce else "manuscript.docx")
    # regenerate tables from results, then number captions (regenerating drops the labels)
    for tool in ("make_tables.py", "number_captions.py"):
        subprocess.run([sys.executable, str(PAPER / "tools" / tool)], check=True, capture_output=True)
    # insert generated tables / blocks at their <!--NAME--> placeholders
    src = (PAPER / "manuscript.md").read_text(encoding="utf-8")
    for f in sorted((PAPER / "tables").glob("*.md")):
        src = src.replace(f"<!--{f.stem.upper()}-->", f.read_text(encoding="utf-8"))
    missing = [m for m in ("TABLE3", "TABLE4", "TABLE5") if f"<!--{m}-->" in src]
    if missing:
        print("warning: placeholders not filled:", missing)
    if asce:
        src = to_asce(src)
    build = PAPER / ("_build_asce.md" if asce else "_build.md")
    build.write_text(src, encoding="utf-8")
    cmd = [PANDOC, build.name, "--citeproc", "--reference-doc", str(REF), "--resource-path", str(PAPER), "-o", str(out)]
    if not asce:
        cmd.insert(3, "--number-sections")
    r = subprocess.run(cmd, cwd=PAPER, capture_output=True, text=True)
    if r.returncode:
        sys.exit(r.stderr)
    if r.stderr.strip():
        print(r.stderr.strip())
    add_line_numbers_and_page_numbers(out)
    hl = PAPER / "highlights.md"
    if hl.exists() and not asce:
        subprocess.run([PANDOC, str(hl), "--reference-doc", str(REF), "-o", str(PAPER / "highlights.docx")],
                       check=True)
    print("->", out)
    if a.pdf:
        export_pdf(out, out.with_suffix(".pdf"))
        print("->", out.with_suffix(".pdf"))


if __name__ == "__main__":
    main()
